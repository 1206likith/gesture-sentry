# -*- coding: utf-8 -*-
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT_PATH = r"L:\Likith\Coding_Projects\ESA_Friend_Project\docs\Sentry_Turret_Full_Wiring_As_Bought.docx"
DIAGRAM = r"L:\Likith\Coding_Projects\ESA_Friend_Project\docs\wiring_diagram.png"


def set_cell_shading(cell, color_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), color_hex)
    tcPr.append(shd)


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = 'Light Grid Accent 1'
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        for p in hdr_cells[i].paragraphs:
            for r in p.runs:
                r.bold = True
                r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        set_cell_shading(hdr_cells[i], "2F5496")
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = str(val)
    return table


def add_bullets(doc, items, style='List Bullet'):
    for it in items:
        doc.add_paragraph(it, style=style)


doc = Document()

# ============================================================
# TITLE
# ============================================================
doc.add_heading('Gesture Sentry Turret \u2014 Full Wiring, As Bought', level=0)
sub = doc.add_paragraph()
sub.add_run(
    'Every connection needed to bring up the turret, matched strictly to the parts on the '
    'three purchase receipts (28-09-2026) \u2014 not the spec\u2019s ideal shopping list.'
).italic = True
meta = doc.add_paragraph()
meta.add_run('Project: ').bold = True
meta.add_run('Gesture-Controlled Sentry Turret (BACSE324, friend\u2019s group)    ')
meta.add_run('Source: ').bold = True
meta.add_run('Gesture_Sentry_Turret_Full_Spec.md + 3 purchase receipts')

doc.add_heading('Parts confirmed on the receipts', level=2)
add_bullets(doc, [
    'Breadboard 400-point, Breadboard 840-point',
    '0.96" OLED display, 4-pin I2C',
    'Female Berg strip 40x1 \u00d7 2',
    '18650 battery holder (4-cell)',
    'Buzzer 5V',
    'Two-cell charger (for 18650 / 16340 / 14500)',
    'MPU6050 (GY-521) 3-axis accelerometer + gyroscope',
    'HC-SR04 ultrasonic sensor',
    'HC-SR501 PIR motion sensor',
    'Laser module',
    'SG90 micro servo',
    'Tower Pro MG995 servo',
    '6x4" general-purpose PCB (perfboard)',
    '18650 rechargeable battery \u00d7 4',
    'ESP32 development board (USB Type-C) \u00d7 2',
    'RGB LED module \u00d7 2',
    'Jumpers \u00d7 16',
    '4-Wheel Drive acrylic chassis kit with motors (separate from this build \u2014 see note at the end)',
])

p = doc.add_paragraph()
p.add_run('Not on any receipt, but required before the build runs untethered: ').bold = True
p.add_run(
    'a buck/boost converter for the turret\u2019s 18650 pack, a LiPo cell + TP4056 + boost converter for '
    'the glove, two resistors for the HC-SR04 voltage divider, and (optional but recommended) a '
    'physical mode-switch button. Each is explained in context below.'
)

doc.add_page_break()

# ============================================================
# DIAGRAM
# ============================================================
doc.add_heading('Full Wiring Diagram', level=1)
doc.add_paragraph(
    'Blue lines are GPIO signal connections, green is the I2C bus to the glove\u2019s IMU, and the '
    'dashed purple line is the BLE wireless link between the two boards \u2014 there is no physical '
    'wire between the turret and the glove.'
)
doc.add_picture(DIAGRAM, width=Inches(7.0))
for p in doc.paragraphs[-1:]:
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_page_break()

# ============================================================
# TURRET ESP32
# ============================================================
doc.add_heading('Turret ESP32 \u2014 Full Connections', level=1)

doc.add_heading('Servos \u2014 MG995 (pan) + SG90 (tilt)', level=2)
doc.add_paragraph(
    'The receipt has an MG995, not the spec\u2019s MG996R \u2014 both share the same 3-wire pinout, so the '
    'wiring is identical either way.'
)
add_table(doc, ['Servo wire', 'Connects to'], [
    ['MG995 (pan) signal \u2014 orange', 'GPIO13'],
    ['MG995 (pan) V+ \u2014 red', 'External 5V rail (NOT the ESP32\u2019s own 5V/USB \u2014 MG995 stall current is too high for that)'],
    ['MG995 (pan) GND \u2014 brown', 'Common GND rail'],
    ['SG90 (tilt) signal \u2014 orange', 'GPIO27 (GPIO12 avoided \u2014 boot-strapping pin on some boards)'],
    ['SG90 (tilt) V+ \u2014 red', 'Same external 5V rail as the pan servo'],
    ['SG90 (tilt) GND \u2014 brown', 'Common GND rail'],
])

doc.add_heading('HC-SR04 ultrasonic', level=2)
add_table(doc, ['Pin', 'Connects to'], [
    ['VCC', '5V rail'],
    ['TRIG', 'GPIO5'],
    ['ECHO', 'GPIO18 \u2014 through a voltage divider (1k\u03a9 + 2k\u03a9); ECHO outputs 5V and the GPIO is 3.3V-only'],
    ['GND', 'Common GND rail'],
])

doc.add_heading('PIR motion sensor (HC-SR501)', level=2)
add_table(doc, ['Pin', 'Connects to'], [
    ['VCC', '5V rail'],
    ['OUT', 'GPIO19 (3.3V logic output, no divider needed)'],
    ['GND', 'Common GND rail'],
])

doc.add_heading('Laser module', level=2)
doc.add_paragraph(
    'The receipt just says "Laser Module," not confirmed as a KY-008. If it is a 3-pin module '
    '(S / + / \u2212) with an onboard resistor, wire as below. If it turns out to be a bare 2-pin '
    'diode with no resistor, do not wire it directly to a GPIO or 3V3 \u2014 add a series resistor '
    '(roughly 100\u2013150\u03a9 for 3.3V) first, or it risks damaging the diode or the pin.'
)
add_table(doc, ['Pin', 'Connects to'], [
    ['S (signal, if 3-pin)', 'GPIO23'],
    ['+ / VCC', '3V3'],
    ['\u2212 / GND', 'Common GND rail'],
])

doc.add_heading('Buzzer (5V)', level=2)
add_table(doc, ['Pin', 'Connects to'], [
    ['+', 'GPIO25'],
    ['\u2212', 'Common GND rail'],
])

doc.add_heading('0.96" OLED (I2C, 4-pin)', level=2)
add_table(doc, ['Pin', 'Connects to'], [
    ['VCC', '3V3'],
    ['GND', 'Common GND rail'],
    ['SCL', 'GPIO22'],
    ['SDA', 'GPIO21'],
])

doc.add_heading('RGB LED module \u00d7 2', level=2)
doc.add_paragraph(
    'The receipt lists two individual "RGB LED Module" units, not the spec\u2019s 12-LED WS2812 ring. '
    'Treat each as a 4-pin common-cathode (or common-anode) RGB LED breakout with separate R/G/B '
    'signal pins \u2014 the simpler and more likely part at this price. One module covers the status '
    'indicator the spec calls for (green = idle, amber = tracking, red = locked); the second is a '
    'spare, the same role the spec\u2019s extra SG90 played.'
)
add_table(doc, ['Pin (common-cathode module)', 'Connects to'], [
    ['R', 'GPIO32 (through the module\u2019s onboard series resistor, or add one if it has none)'],
    ['G', 'GPIO33'],
    ['B', 'GPIO4'],
    ['GND (common pin)', 'Common GND rail'],
])
doc.add_paragraph(
    'If the module is common-anode instead, wire the common pin to 3V3 instead of GND, and drive '
    'R/G/B LOW to turn each color on (inverted logic) in firmware.'
)

doc.add_heading('Mode-switch button', level=2)
doc.add_paragraph(
    'Not on the receipt. If a spare push-button exists, wire one leg to GPIO26 and the other to '
    'GND (use INPUT_PULLUP in firmware, no external resistor needed). Without one, the mode-switch '
    'gesture (double-tap) is the only way to flip AUTO/MANUAL \u2014 fine for bench testing, but the '
    'spec recommends a physical fallback before the live demo in case that gesture misclassifies.'
)

doc.add_heading('Turret power path', level=2)
doc.add_paragraph(
    '4\u00d7 18650 cells go into the 4-cell holder, but the pack\u2019s wiring (series or parallel) is not '
    'yet confirmed \u2014 check it with a multimeter before connecting anything. This matters because it '
    'determines whether a boost converter (parallel pack, ~3.7V) or a buck converter (series pack, '
    '7.4\u201314.8V+) is needed to get a clean regulated 5V. Neither is on any receipt.'
)
doc.add_paragraph(
    'The receipt\u2019s "two-cell charger" only charges cells outside the circuit \u2014 it is not inline '
    'power regulation and does not belong in the live wiring path.'
)
p = doc.add_paragraph()
p.add_run('Until a matched buck/boost module is bought and the pack voltage is confirmed: ').bold = True
p.add_run('run the turret on USB power for all bring-up and testing.')

doc.add_heading('Turret ESP32 pin summary', level=2)
add_table(doc, ['GPIO', 'Function'], [
    ['4', 'RGB LED module \u2014 Blue'],
    ['5', 'HC-SR04 TRIG'],
    ['13', 'MG995 pan signal'],
    ['18', 'HC-SR04 ECHO (via divider)'],
    ['19', 'PIR OUT'],
    ['21', 'OLED SDA'],
    ['22', 'OLED SCL'],
    ['23', 'Laser signal'],
    ['25', 'Buzzer'],
    ['26', 'Mode-switch button (if present)'],
    ['27', 'SG90 tilt signal'],
    ['32', 'RGB LED module \u2014 Red'],
    ['33', 'RGB LED module \u2014 Green'],
    ['BLE (internal radio)', 'Receives gesture labels from the Glove ESP32'],
])

doc.add_page_break()

# ============================================================
# GLOVE ESP32
# ============================================================
doc.add_heading('Glove ESP32 \u2014 Full Connections', level=1)

doc.add_heading('MPU6050 IMU (I2C)', level=2)
add_table(doc, ['Pin', 'Connects to'], [
    ['VCC', '3V3'],
    ['GND', 'Common GND rail'],
    ['SCL', 'GPIO22'],
    ['SDA', 'GPIO21'],
    ['AD0', 'GND (sets I2C address 0x68, the default)'],
    ['INT', 'Not used \u2014 polling mode is fine for this model'],
])

doc.add_heading('Glove power path \u2014 gap vs. the spec', level=2)
p = doc.add_paragraph()
p.add_run('No LiPo cell or TP4056 is on any receipt. ').bold = True
p.add_run(
    'The glove currently has no portable power source. Until one is bought, it runs tethered via '
    'USB cable to a laptop or power bank \u2014 this defeats the "untethered glove" part of the demo but '
    'does not block firmware or gesture-training work, all of which can be done over USB. To go '
    'wireless: buy a small LiPo pouch cell, a TP4056 charge module, and a boost converter (three '
    'items, none currently owned) \u2014 the same power chain the spec describes for the glove.'
)

doc.add_heading('Glove ESP32 pin summary', level=2)
add_table(doc, ['GPIO', 'Function'], [
    ['21', 'MPU6050 SDA'],
    ['22', 'MPU6050 SCL'],
    ['BLE (internal radio)', 'Transmits classified gesture label to the Turret ESP32'],
])

doc.add_heading('Cross-board link', level=2)
doc.add_paragraph(
    'There is no wired connection between the two ESP32s \u2014 they communicate only over BLE '
    '(Glove = peripheral broadcasting gesture labels, Turret = central receiving them). This is the '
    'one link in the whole system that is wireless by design, per the spec\u2019s control architecture.'
)

doc.add_page_break()

# ============================================================
# PROTOTYPING + LEFTOVER ITEMS
# ============================================================
doc.add_heading('Prototyping-Stage Parts (bought, used before final assembly)', level=1)
add_bullets(doc, [
    'Breadboard 400-point + 840-point \u2014 bench-test every module above before soldering to the '
    'perfboard. Put the turret\u2019s modules on the 840-point (more room for servos, sensors, OLED, '
    'and LEDs together); the MPU6050 alone fits the 400-point for the glove side.',
    'Female Berg strip \u00d7 2 \u2014 solder onto the perfboard as socket headers so the ESP32s and '
    'modules can be plugged/unplugged from the final board rather than soldered directly \u2014 standard '
    'practice, do this before final assembly.',
    'Jumpers \u00d7 16 \u2014 enough for bringing up one subsystem at a time on a breadboard, but tight for '
    'the full turret wired simultaneously (servos + ultrasonic + PIR + laser + buzzer + OLED + RGB '
    'LED is 20+ individual wires). Budget for more jumpers, or move to perfboard sooner.',
])

doc.add_heading('4WD Acrylic Chassis Kit (\u20b9649) \u2014 Not Part of This Build', level=1)
doc.add_paragraph(
    'This is a wheeled-robot chassis (motors + wheels + acrylic deck) matching a different '
    'brainstormed idea \u2014 a line-follower/obstacle-avoider or the sumo robot \u2014 not the sentry turret '
    'spec. No motor driver (L298N or similar) was bought to go with it, so it cannot drive anything '
    'yet. Keep it aside unless the team is also building a second, separate robot.'
)

doc.add_heading('Summary \u2014 What to Buy Before the Build Is Fully Untethered', level=1)
add_bullets(doc, [
    'Buck or boost converter for the turret\u2019s 18650 pack (confirm series vs. parallel wiring first '
    'with a multimeter \u2014 this decides which type is needed).',
    'LiPo pouch cell + TP4056 charge module + boost converter for the glove.',
    'Two resistors (1k\u03a9 + 2k\u03a9) for the HC-SR04 ECHO voltage divider.',
    'A small push-button for the physical mode-switch fallback (optional but recommended).',
    'An L298N motor driver \u2014 only if the 4WD chassis is used on a separate build.',
])

doc.save(OUT_PATH)
print("Saved:", OUT_PATH)
print("Paragraphs:", len(doc.paragraphs))
print("Tables:", len(doc.tables))
